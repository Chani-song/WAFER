# -*- coding: utf-8 -*-
"""Experimental LibreOffice CLI backend for BatPaC.

Use only as a proof of concept until validated against Microsoft Excel.

Flow:
1. Copy workbook to a temporary runtime directory.
2. Write Dashboard override cells with openpyxl.
3. Ask LibreOffice headless to open/recalculate/save as xlsx.
4. Read Dashboard!B187:J214 and optional cell_outputs with openpyxl.

This module intentionally exposes the same functions used by src.tools.batpac:
- batpac_by_python(...)
- _parse_numeric(...)
"""

from __future__ import annotations

import ast
import datetime
import math
import os
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path
from typing import Any, Dict, Union

import pandas as pd
from openpyxl import load_workbook


_CELL_RE = re.compile(r"^[A-Z]+[0-9]+$")


def _project_root() -> Path:
    return Path(__file__).resolve().parent


def _load_cell_address_map() -> Dict[str, str]:
    """Read CELL_ADDRESS_MAP from batpac_codegen2.py without importing xlwings."""
    codegen_path = _project_root() / "batpac_codegen2.py"
    tree = ast.parse(codegen_path.read_text(encoding="utf-8"))
    for node in tree.body:
        if isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name):
            if node.target.id == "CELL_ADDRESS_MAP" and node.value is not None:
                return ast.literal_eval(node.value)
        if isinstance(node, ast.Assign):
            for target in node.targets:
                if isinstance(target, ast.Name) and target.id == "CELL_ADDRESS_MAP":
                    return ast.literal_eval(node.value)
    raise RuntimeError(f"Could not find CELL_ADDRESS_MAP in {codegen_path}")


def _kwargs_to_addrvals(overrides: Dict[str, Any]) -> Dict[str, Any]:
    cell_map = _load_cell_address_map()
    addrvals: Dict[str, Any] = {}

    for key, value in overrides.items():
        if key in cell_map:
            addrvals[cell_map[key]] = value

    for key, value in overrides.items():
        key_text = str(key).upper().strip()
        if _CELL_RE.fullmatch(key_text):
            addrvals[key_text] = value

    return addrvals


def _parse_numeric(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return value

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("$", "").replace(",", "").replace("%", "")
    try:
        return float(text)
    except Exception:
        return value


def _write_dashboard_openpyxl(
    xlsm_path: Path,
    sheet: str,
    addrvals: Dict[str, Any],
    safe_only: bool = True,
) -> None:
    wb = load_workbook(xlsm_path, keep_vba=True)
    try:
        ws = wb[sheet]
        for addr, val in addrvals.items():
            a = str(addr).upper().strip()
            m = re.fullmatch(r"([A-Z]+)([0-9]+)", a)
            if not m:
                continue
            col, row = m.group(1), int(m.group(2))
            if safe_only and (row < 1 or row > 400 or col > "L"):
                continue
            ws[a] = val
        wb.save(xlsm_path)
    finally:
        wb.close()


def _run_libreoffice_convert(input_path: Path, out_dir: Path, timeout_s: int = 240) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    cmd = [
        "libreoffice",
        "--headless",
    ]
    user_profile = os.environ.get("BATPAC_LO_USER_PROFILE")
    if user_profile:
        profile_path = Path(user_profile).resolve()
        profile_path.mkdir(parents=True, exist_ok=True)
        cmd.append(f"-env:UserInstallation={profile_path.as_uri()}")
    cmd.extend([
        "--convert-to",
        "xlsx",
        "--outdir",
        str(out_dir),
        str(input_path),
    ])
    proc = subprocess.run(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        timeout=timeout_s,
    )

    if proc.returncode != 0:
        raise RuntimeError("LibreOffice conversion failed:\n" + proc.stdout)

    out_xlsx = out_dir / (input_path.stem + ".xlsx")
    if not out_xlsx.exists():
        raise FileNotFoundError(
            f"LibreOffice did not create expected file: {out_xlsx}\nOutput:\n{proc.stdout}"
        )
    return out_xlsx


def _read_range(xlsx_path: Path, sheet: str, cell_range: str):
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        ws = wb[sheet]
        return [[cell.value for cell in row] for row in ws[cell_range]]
    finally:
        wb.close()


def get_results_block_from_xlsx(
    xlsx_path: Union[str, Path],
    result_sheet: str = "Dashboard",
    header_row: int = 187,
    data_start: int = 188,
    data_end: int = 214,
    start_col: str = "B",
    end_col: str = "J",
):
    xlsx_path = Path(xlsx_path)
    headers = _read_range(xlsx_path, result_sheet, f"{start_col}{header_row}:{end_col}{header_row}")[0]
    data = _read_range(xlsx_path, result_sheet, f"{start_col}{data_start}:{end_col}{data_end}")

    if not headers or all(h is None or str(h).strip() == "" for h in headers):
        n_cols = ord(end_col) - ord(start_col) + 1
        headers = ["Index", "Attribute"] + [f"Battery {i}" for i in range(1, n_cols - 1)]

    df = pd.DataFrame(data, columns=[str(h).strip() if h is not None else "" for h in headers])
    if df.columns.size >= 2:
        df.rename(columns={df.columns[0]: "Index", df.columns[1]: "Attribute"}, inplace=True)
    return df


def get_named_cell_values_from_xlsx(
    xlsx_path: Union[str, Path],
    cell_specs: Dict[str, Dict[str, str]],
) -> Dict[str, Any]:
    xlsx_path = Path(xlsx_path)
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    values: Dict[str, Any] = {}
    try:
        for name, spec in (cell_specs or {}).items():
            sheet = str(spec["sheet"])
            cell = str(spec["cell"]).upper()
            values[name] = wb[sheet][cell].value
    finally:
        wb.close()
    return values


def batpac_by_python(
    xlsm_path: Union[str, Path],
    dashboard_sheet: str = "Dashboard",
    result_sheet: str = "Dashboard",
    visible: bool = False,
    out_csv: Union[None, str, Path, bool] = None,
    return_timing: bool = True,
    extra_cells: Union[None, Dict[str, Dict[str, str]]] = None,
    **kwargs,
):
    del visible  # ignored by LibreOffice backend

    xlsm_path = Path(xlsm_path)
    t0 = time.perf_counter()

    addrvals = _kwargs_to_addrvals(kwargs)
    # Match the current Excel runner's "hold" behavior.
    # Dashboard!D6 is a restart/reconvergence trigger.
    # LibreOffice cannot use the Excel macro/Ctrl+Shift+R path, so this backend
    # reads outputs in the post-restart/off state with D6=0, matching Excel pulse/off.
    addrvals["D6"] = 0
    t_inputs = time.perf_counter()

    with tempfile.TemporaryDirectory(prefix="batpac_lo_") as tmp:
        tmp_dir = Path(tmp)
        runtime_xlsm = tmp_dir / xlsm_path.name
        shutil.copy2(xlsm_path, runtime_xlsm)

        t_open0 = time.perf_counter()
        _write_dashboard_openpyxl(runtime_xlsm, dashboard_sheet, addrvals, safe_only=True)
        t_write = time.perf_counter()

        runtime_xlsx = _run_libreoffice_convert(runtime_xlsm, tmp_dir)

        # Optional debug: keep the LibreOffice-calculated xlsx for backend validation.
        # Enable with BATPAC_KEEP_LO_XLSX=1. Normal runs should not accumulate xlsx files.
        if os.environ.get("BATPAC_KEEP_LO_XLSX", "").strip() == "1":
            debug_dir = _project_root() / "outputs_mcp" / "backend_validation" / "lo_debug_xlsx"
            debug_dir.mkdir(parents=True, exist_ok=True)
            debug_xlsx = debug_dir / f"{runtime_xlsm.stem}_calculated.xlsx"
            shutil.copy2(runtime_xlsx, debug_xlsx)

        t_recalc = time.perf_counter()

        df = get_results_block_from_xlsx(
            runtime_xlsx,
            result_sheet=result_sheet,
            header_row=187,
            data_start=188,
            data_end=214,
            start_col="B",
            end_col="J",
        )
        t_read = time.perf_counter()

        extra_values = get_named_cell_values_from_xlsx(runtime_xlsx, extra_cells or {})
        t_extra = time.perf_counter()

        if out_csv is not None:
            if out_csv is True:
                ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                out_path = xlsm_path.with_suffix(f".LO_B187_J214_{ts}.csv")
            else:
                out_path = Path(out_csv)
            df.to_csv(out_path, index=False)

        t1 = time.perf_counter()

    if return_timing:
        timings = {
            "build_inputs_s": t_inputs - t0,
            "excel_open_s": 0.0,
            "write_inputs_s": t_write - t_open0,
            "recalc_s": t_recalc - t_write,
            "read_results_s": t_read - t_recalc,
            "read_extra_cells_s": t_extra - t_read,
            "total_s": t1 - t0,
            "backend": "libreoffice_cli",
        }
        if extra_cells is not None:
            return df, timings, extra_values
        return df, timings

    if extra_cells is not None:
        return df, extra_values
    return df
