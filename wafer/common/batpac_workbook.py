"""Read-only checks and reads on the user's own BatPaC workbook (release addition, no workbook values stored)."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any, Dict, Mapping, Optional, Union

REQUIRED_VERSION = "5.0"
EXPECTED_FILE_NAME = "BatPaC 5.0 2022-02-09.xlsm"


def batpac_workbook_version(xlsm_path: Union[str, Path]) -> Optional[str]:
    """Return the version printed on the About sheet ("BatPaC-Version 5.0"), or None if absent."""
    with zipfile.ZipFile(xlsm_path) as z:
        for name in z.namelist():
            if name.startswith("xl/drawings/") and name.endswith(".xml"):
                m = re.search(rb"BatPaC-Version\s*([0-9][0-9.]*)", z.read(name))
                if m:
                    return m.group(1).decode()
    return None


def require_batpac_5(xlsm_path: Union[str, Path]) -> str:
    """Stop with a clear message unless the workbook exists and reports BatPaC version 5.0."""
    path = Path(xlsm_path)
    if not path.is_file():
        raise FileNotFoundError(
            f"BatPaC workbook not found: {path}. Obtain BatPaC 5.0 (2022-02-09) from Argonne National "
            f"Laboratory and pass it with --batpac-workbook (see README, 'BatPaC').")
    try:
        version = batpac_workbook_version(path)
    except zipfile.BadZipFile as exc:
        raise RuntimeError(f"{path} is not an .xlsm workbook: {exc}") from exc
    if version != REQUIRED_VERSION:
        raise RuntimeError(
            f"{path} reports BatPaC version {version!r}; this code requires BatPaC {REQUIRED_VERSION} "
            f"(2022-02-09). Cell addresses differ between BatPaC versions.")
    return version


def read_workbook_defaults(xlsm_path: Union[str, Path], cell_address_map: Mapping[str, str],
                           sheet: str = "Dashboard") -> Dict[str, Any]:
    """Current (cached) value of every mapped Dashboard cell in the user's workbook."""
    require_batpac_5(xlsm_path)
    from openpyxl import load_workbook

    ws = load_workbook(xlsm_path, data_only=True)[sheet]
    return {name: ws[address].value for name, address in cell_address_map.items()}
